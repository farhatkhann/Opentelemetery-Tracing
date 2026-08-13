import json
from pathlib import Path
from statistics import mean, stdev
from openpyxl import Workbook, load_workbook


# ==========================================================
# Configuration
# ==========================================================

APPLICATION = "Spring PetClinic"

BASE_DIR = Path("../datasets/traces/monolith")

SAMPLINGS = {
    "off": "Off",
    "sample100": "100%",
    "sample10": "10%",
    "sample1": "1%"
}

RUNS = 5

OUTPUT_FILE = BASE_DIR / "spring_petclinic_results.xlsx"


# ==========================================================
# Create / Open Workbook
# ==========================================================

if OUTPUT_FILE.exists():
    wb = load_workbook(OUTPUT_FILE)

    if "Results" in wb.sheetnames:
        ws = wb["Results"]
    else:
        ws = wb.create_sheet("Results")

else:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"


# ==========================================================
# Create Header if Sheet is Empty
# ==========================================================

headers = [
    "Application",
    "Service",
    "Workflow",
    "Sampling",

    "Avg Response Time (ms)",
    "SD",

    "Throughput (req/s)",
    "SD",

    "Trace Count",
    "SD",

    "Span Count",
    "SD",

    "Unique Operations",
    "SD",

    "Trace File Size (KB)",
    "SD",

    "Avg Span Duration (µs)",
    "SD",

    "CPU Utilization (%)",
    "SD",

    "Memory Usage (MB)",
    "SD"
]

if ws.max_row == 1 and ws.cell(1, 1).value is None:
    ws.append(headers)

elif ws.max_row == 0:
    ws.append(headers)


# ==========================================================
# Helper Functions
# ==========================================================

def get_existing_row(service, workflow, sampling):

    for row in range(2, ws.max_row + 1):

        if (
            ws.cell(row, 2).value == service
            and ws.cell(row, 3).value == workflow
            and ws.cell(row, 4).value == sampling
        ):
            return row

    return None


def safe_mean(values):

    if not values:
        return 0

    return round(mean(values), 2)


def safe_sd(values):

    if len(values) <= 1:
        return 0

    return round(stdev(values), 2)


# ==========================================================
# Process Each Sampling
# ==========================================================

for sampling_folder, sampling_display in SAMPLINGS.items():

    sampling_path = BASE_DIR / sampling_folder

    print("\n" + "=" * 70)
    print(f"Processing Sampling: {sampling_display}")
    print(f"Folder: {sampling_path}")
    print("=" * 70)

    if not sampling_path.exists():

        print(f"WARNING: Sampling folder does not exist:")
        print(f"        {sampling_path}")

        continue


    # ======================================================
    # Find Workflows
    # ======================================================

    workflow_folders = [
        folder
        for folder in sampling_path.iterdir()
        if folder.is_dir()
    ]

    if not workflow_folders:

        print(f"No workflow folders found in {sampling_path}")
        continue


    # ======================================================
    # Process Each Workflow
    # ======================================================

    for workflow_path in sorted(workflow_folders):

        workflow = workflow_path.name

        print(f"\nWorkflow: {workflow}")


        # Lists for multiple runs
        response_times = []
        throughputs = []
        trace_counts = []
        span_counts = []
        unique_operations = []
        trace_sizes = []
        avg_span_durations = []
        cpu_utils = []
        memory_utils = []

        service = None


        # ==================================================
        # Process Runs
        # ==================================================

        for run in range(1, RUNS + 1):

            run_path = workflow_path / f"run{run}"

            experiment_file = run_path / "experiment.json"
            summary_file = run_path / "summary.json"


            if not run_path.exists():

                print(f"  Missing run folder: {run_path}")

                continue


            if not experiment_file.exists():

                print(f"  Missing: {experiment_file}")

                continue


            if not summary_file.exists():

                print(f"  Missing: {summary_file}")

                continue


            print(f"  Processing run{run}")


            # ----------------------------------------------
            # Read JSON files
            # ----------------------------------------------

            try:

                with open(experiment_file, "r", encoding="utf-8") as f:
                    experiment = json.load(f)


                with open(summary_file, "r", encoding="utf-8") as f:
                    summary = json.load(f)

            except Exception as e:

                print(f"  ERROR reading JSON files: {e}")

                continue


            # ----------------------------------------------
            # Service
            # ----------------------------------------------

            service = experiment.get("service", "Unknown")


            # ----------------------------------------------
            # Performance Metrics
            # ----------------------------------------------

            performance = experiment.get(
                "performance_metrics",
                {}
            )

            resource = experiment.get(
                "resource_metrics",
                {}
            )


            if "average_response_time_ms" in performance:

                response_times.append(
                    performance["average_response_time_ms"]
                )


            if "throughput_req_per_sec" in performance:

                throughputs.append(
                    performance["throughput_req_per_sec"]
                )


            # ----------------------------------------------
            # Resource Metrics
            # ----------------------------------------------

            if "cpu_utilization_percent" in resource:

                cpu_utils.append(
                    resource["cpu_utilization_percent"]
                )


            if "memory_usage_mb" in resource:

                memory_utils.append(
                    resource["memory_usage_mb"]
                )


            # ----------------------------------------------
            # Trace Metrics
            # ----------------------------------------------

            if "trace_count" in summary:

                trace_counts.append(
                    summary["trace_count"]
                )


            if "span_count" in summary:

                span_counts.append(
                    summary["span_count"]
                )


            if "unique_operations" in summary:

                unique_operations.append(
                    summary["unique_operations"]
                )


            if "trace_file_size_kb" in summary:

                trace_sizes.append(
                    summary["trace_file_size_kb"]
                )


            if "average_span_duration_us" in summary:

                avg_span_durations.append(
                    summary["average_span_duration_us"]
                )


        # ==================================================
        # No Valid Runs
        # ==================================================

        if service is None:

            print(
                f"  No valid runs found for "
                f"{workflow} - {sampling_display}"
            )

            continue


        # ==================================================
        # Create Excel Row
        # ==================================================

        row = [

            APPLICATION,

            service,

            workflow,

            sampling_display,


            # Response Time
            safe_mean(response_times),
            safe_sd(response_times),


            # Throughput
            safe_mean(throughputs),
            safe_sd(throughputs),


            # Trace Count
            safe_mean(trace_counts),
            safe_sd(trace_counts),


            # Span Count
            safe_mean(span_counts),
            safe_sd(span_counts),


            # Unique Operations
            safe_mean(unique_operations),
            safe_sd(unique_operations),


            # Trace File Size
            safe_mean(trace_sizes),
            safe_sd(trace_sizes),


            # Average Span Duration
            safe_mean(avg_span_durations),
            safe_sd(avg_span_durations),


            # CPU
            safe_mean(cpu_utils),
            safe_sd(cpu_utils),


            # Memory
            safe_mean(memory_utils),
            safe_sd(memory_utils)

        ]


        # ==================================================
        # Update Existing Row
        # ==================================================

        existing_row = get_existing_row(
            service,
            workflow,
            sampling_display
        )


        if existing_row:

            for col, value in enumerate(row, start=1):

                ws.cell(
                    existing_row,
                    col
                ).value = value

            print(
                f"  Updated Excel row: "
                f"{workflow} - {sampling_display}"
            )


        # ==================================================
        # Add New Row
        # ==================================================

        else:

            ws.append(row)

            print(
                f"  Added Excel row: "
                f"{workflow} - {sampling_display}"
            )


# ==========================================================
# Sort Excel Rows
# ==========================================================

# Desired sampling order
sampling_order = {
    "Off": 0,
    "100%": 1,
    "10%": 2,
    "1%": 3
}


# Read all data rows
data = list(
    ws.iter_rows(
        min_row=2,
        values_only=True
    )
)


# Sort by:
# 1. Workflow
# 2. Sampling
data.sort(
    key=lambda row: (
        row[2] if row[2] is not None else "",
        sampling_order.get(row[3], 99)
    )
)


# Delete existing data rows
if ws.max_row > 1:

    ws.delete_rows(
        2,
        ws.max_row - 1
    )


# Add sorted rows
for row in data:

    ws.append(list(row))


# ==========================================================
# Auto Adjust Column Width
# ==========================================================

for column_cells in ws.columns:

    max_length = 0

    for cell in column_cells:

        if cell.value is not None:

            max_length = max(
                max_length,
                len(str(cell.value))
            )


    column_letter = column_cells[0].column_letter

    ws.column_dimensions[
        column_letter
    ].width = max_length + 3


# ==========================================================
# Save Workbook
# ==========================================================

try:

    wb.save(OUTPUT_FILE)

    print("\n")
    print("=" * 70)
    print("RESULTS GENERATED SUCCESSFULLY")
    print("=" * 70)
    print(f"Output File: {OUTPUT_FILE}")
    print("=" * 70)


except PermissionError:

    print("\nERROR: Cannot save the Excel file.")

    print(
        "Please close "
        "'spring_petclinic_results.xlsx' "
        "if it is open in Excel."
    )