#!/usr/bin/env python3
"""
aggregate_coverage_by_sampling.py
------------------------------------
Computes the coverage metric that actually answers "how much of the app's
endpoint surface does this sampling ratio capture?": for a given sampling
ratio, union the observed (service, http_method, route) triples across
EVERY workflow's every run, then compare that union against the full
static inventory (extract_static_inventory_java.py's static_inventory.csv).

This is different from (and more meaningful than) averaging per-workflow
coverage numbers -- a single k6 workflow only ever exercises a subset of
this app's endpoints, so per-workflow coverage is capped well below 100%
even with perfect tracing. Aggregating across all workflows first removes
that artifact -- this is the number that actually answers "how much
coverage does this sampling ratio give me system-wide".

Expects the SAME folder layout collect_traces.py produces and that
run_quality_analysis.py reads (edit SAMPLINGS below if your folder names
differ):
    {base-dir}/{sampling}/{workflow}/run{N}/traces.json

Writes both:
  - a CSV file with one row per sampling ratio (--out)
  - an Excel workbook (by default saved into --base-dir, alongside
    run_quality_analysis.py's quality_summary.xlsx) with three sheets:
      1. Summary            -- one row per sampling ratio: trace/span
                                totals, endpoints matched/missing, coverage %
      2. Endpoint Matrix    -- one row per endpoint, one column per
                                sampling ratio, Yes/No whether it was
                                observed at that ratio
      3. Per-Workflow Hits  -- one row per (sampling, workflow, endpoint)
                                actually observed, so you can see which
                                workflow is responsible for covering which
                                endpoint

Usage:
    python aggregate_coverage_by_sampling.py \
        --base-dir ../datasets/traces/microservices \
        --inventory static_inventory.csv \
        --out aggregate_coverage_by_sampling.csv
    # Excel defaults to {base-dir}/aggregate_coverage_by_sampling.xlsx
    # Pass --excel-out to override.
"""
import argparse
import csv
import os
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from compute_coverage import load_inventory, build_observed_index

SAMPLINGS = {
    "off": "Off",
    "sample1": "1%",
    "sample10": "10%",
    "sample100": "100%",
}


def endpoint_display(entry):
    return f"[{entry['service']}] {entry['http_method'].upper()} {entry['route_template_normalized']}"


def endpoint_key(entry):
    return (entry["service"], entry["http_method"].upper(), entry["route_template_normalized"])


def find_run_traces(base_dir: Path):
    """
    Yields (sampling_folder, workflow_name, run_number, traces_path) for
    every traces.json found under base_dir, in the
    {sampling}/{workflow}/run{N}/traces.json layout. Same layout/logic as
    run_quality_analysis.py's find_run_traces, duplicated here so this
    script can run standalone.
    """
    for sampling_dir in sorted(base_dir.iterdir()):
        if not sampling_dir.is_dir() or sampling_dir.name not in SAMPLINGS:
            continue
        for workflow_dir in sorted(sampling_dir.iterdir()):
            if not workflow_dir.is_dir():
                continue
            for run_dir in sorted(workflow_dir.iterdir()):
                if not run_dir.is_dir() or not run_dir.name.startswith("run"):
                    continue
                traces_path = run_dir / "traces.json"
                if traces_path.exists():
                    run_num = int(run_dir.name.replace("run", ""))
                    yield sampling_dir.name, workflow_dir.name, run_num, traces_path


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--base-dir", required=True, help="datasets/traces/microservices")
    ap.add_argument("--inventory", required=True, help="static_inventory.csv")
    ap.add_argument("--out", default="aggregate_coverage_by_sampling.csv")
    ap.add_argument("--excel-out", default=None,
                     help="Default: {base-dir}/aggregate_coverage_by_sampling.xlsx")
    args = ap.parse_args()

    base_dir = Path(args.base_dir)
    inventory = load_inventory(args.inventory)
    total_endpoints = len(inventory)

    runs = list(find_run_traces(base_dir))
    if not runs:
        print(f"No traces.json files found under {base_dir} matching "
              f"{{sampling}}/{{workflow}}/run{{N}}/traces.json")
        return

    results = {}  # sampling -> dict of aggregate data

    for sampling in SAMPLINGS:
        entries = [(wf, run, path) for (s, wf, run, path) in runs if s == sampling]
        if not entries:
            print(f"{sampling}: no traces.json files found, skipping")
            continue

        union_observed = {}      # (service, METHOD, route) -> total occurrence count
        total_traces = 0
        total_spans = 0
        per_workflow_hits = {}   # workflow -> set of endpoint_display() strings

        for workflow, run_num, traces_path in entries:
            observed, n_traces, n_spans, _ = build_observed_index(str(traces_path))
            total_traces += n_traces
            total_spans += n_spans
            for key, count in observed.items():
                union_observed[key] = union_observed.get(key, 0) + count

            hit_here = {
                endpoint_display(entry)
                for entry in inventory
                if endpoint_key(entry) in observed
            }
            per_workflow_hits.setdefault(workflow, set()).update(hit_here)

        matched = sorted(endpoint_display(e) for e in inventory if endpoint_key(e) in union_observed)
        missing = sorted(endpoint_display(e) for e in inventory if endpoint_key(e) not in union_observed)
        coverage_pct = round(100 * len(matched) / total_endpoints, 2) if total_endpoints else 0.0

        results[sampling] = {
            "runs_included": [(wf, run) for wf, run, _ in entries],
            "total_traces": total_traces,
            "total_spans": total_spans,
            "route_coverage_percent": coverage_pct,
            "matched_routes": matched,
            "missing_routes": missing,
            "per_workflow_route_hits": {wf: sorted(hits) for wf, hits in per_workflow_hits.items()},
        }

        print(f"{sampling}: {coverage_pct}% ({len(matched)}/{total_endpoints} endpoints) "
              f"across {len(entries)} workflow-run(s), {total_traces} traces")

    # ---- CSV summary ----
    fieldnames = ["sampling", "sampling_label", "workflow_runs_included", "total_traces", "total_spans",
                  "endpoints_matched", "endpoints_missing", "endpoints_total", "endpoint_coverage_pct"]
    with open(args.out, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for sampling in SAMPLINGS:
            if sampling not in results:
                continue
            r = results[sampling]
            writer.writerow({
                "sampling": sampling,
                "sampling_label": SAMPLINGS[sampling],
                "workflow_runs_included": len(r["runs_included"]),
                "total_traces": r["total_traces"],
                "total_spans": r["total_spans"],
                "endpoints_matched": len(r["matched_routes"]),
                "endpoints_missing": len(r["missing_routes"]),
                "endpoints_total": total_endpoints,
                "endpoint_coverage_pct": r["route_coverage_percent"],
            })
    print(f"\nSaved CSV summary to {args.out}")

    excel_path = Path(args.excel_out) if args.excel_out else base_dir / "aggregate_coverage_by_sampling.xlsx"
    write_excel(results, inventory, excel_path)
    print(f"Saved Excel summary to {excel_path}")

    if "off" in results and results["off"]["route_coverage_percent"] > 0:
        print("\nNOTE: 'off' sampling shows nonzero coverage -- check whether tracing "
              "was actually disabled for that capture, or spans are being emitted "
              "regardless of the sampler setting.")


def autosize_columns(ws):
    for column_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(length + 3, 60)


def write_excel(results, inventory, excel_path: Path):
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    yes_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    no_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    def style_header(ws, row=1):
        for cell in ws[row]:
            cell.font = header_font
            cell.fill = header_fill

    present_samplings = [s for s in SAMPLINGS if s in results]

    # ---- Sheet 1: Summary ----
    ws = wb.active
    ws.title = "Summary"
    ws.append([
        "Sampling Ratio", "Total Traces", "Total Spans",
        "Endpoints Matched", "Endpoints Missing", "Total Endpoints",
        "Endpoint Coverage %",
    ])
    for sampling in present_samplings:
        r = results[sampling]
        ws.append([
            SAMPLINGS.get(sampling, sampling),
            r["total_traces"],
            r["total_spans"],
            len(r["matched_routes"]),
            len(r["missing_routes"]),
            len(r["matched_routes"]) + len(r["missing_routes"]),
            r["route_coverage_percent"],
        ])
    style_header(ws)
    autosize_columns(ws)

    # ---- Sheet 2: Endpoint Matrix (endpoint x sampling ratio) ----
    ws2 = wb.create_sheet("Endpoint Matrix")
    ws2.append(["Endpoint"] + [SAMPLINGS.get(s, s) for s in present_samplings])
    for entry in inventory:
        display = endpoint_display(entry)
        row = [display]
        for sampling in present_samplings:
            hit = display in results[sampling]["matched_routes"]
            row.append("Yes" if hit else "No")
        ws2.append(row)
    style_header(ws2)
    for row in ws2.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.fill = yes_fill if cell.value == "Yes" else no_fill
    autosize_columns(ws2)

    # ---- Sheet 3: Per-Workflow Hits ----
    ws3 = wb.create_sheet("Per-Workflow Hits")
    ws3.append(["Sampling Ratio", "Workflow", "Endpoint Hit"])
    for sampling in present_samplings:
        for workflow, hit_routes in sorted(results[sampling]["per_workflow_route_hits"].items()):
            for route in hit_routes:
                ws3.append([SAMPLINGS.get(sampling, sampling), workflow, route])
    style_header(ws3)
    autosize_columns(ws3)

    wb.save(excel_path)


if __name__ == "__main__":
    main()
